"""검증 워크북 (RESULT.md 완료 기준 1).

독립성 설계
- 파이썬은 TTM을 '분기값 4개 합산'으로 계산한다.
- 워크북은 엑셀 수식으로 'YTD(당기) + 연간(전년) − YTD(전년 동기)'로 다시 계산한다 (다른 경로).
- 원천값(계정 매핑 결과)은 사람이 DART 원문 링크로 직접 대조한다 → '원천값 확인' 시트의 확인 칸.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FLOWS = [("ni_owners", "지배주주순이익"), ("ni_total", "당기순이익"), ("revenue", "매출액"), ("gross_profit", "매출총이익"),
         ("op_income", "영업이익"), ("cfo", "영업현금흐름")]
STOCKS = [("equity_owners", "지배주주자본"), ("assets", "자산총계"), ("debt", "이자부부채"), ("cash", "현금성자산"),
          ("st_fin", "단기금융상품")]
HDR = PatternFill("solid", fgColor="DDE7F7")
INPUT = PatternFill("solid", fgColor="FFF6D6")
EXCEL = PatternFill("solid", fgColor="E8F5E9")
TOL = "1E-9"


def _match(a: str, b: str) -> str:
    return (f'=IF(AND({a}="",{b}=""),"",IF(OR({a}="",{b}=""),"불일치",'
            f'IF(ABS({a}-{b})<=MAX(1,ABS({b}))*{TOL},"일치","불일치")))')


def build(results: list[dict], path: Path) -> None:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "안내"
    for i, line in enumerate([
        "1단계 계산 검증 워크북 (자동 생성)",
        "",
        "■ 검증계산 시트",
        "  · 노란 칸 = 파이썬이 DART에서 뽑은 원천값 (YTD·잔액)",
        "  · 초록 칸 = 이 엑셀 수식으로 다시 계산한 값. TTM은 'YTD(당기) + 연간(전년) − YTD(전년 동기)' 방식",
        "    (파이썬은 '분기값 4개 합산' 방식 → 서로 다른 경로로 같은 답이 나와야 함)",
        "  · '일치/불일치' 칸이 전부 '일치'여야 계산 로직이 맞는 것",
        "",
        "■ 원천값 확인 시트 (사람이 확인)",
        "  · 각 행의 DART 링크를 열어 해당 계정 금액이 맞는지 확인 후 '확인' 칸에 O/X 입력",
        "  · X가 나오면 계정 매핑 규칙(config/account_rules.yaml) 수정 대상",
        "",
        "■ 커버리지 시트: 지표별 계산 가능 비율과 무효 사유",
        "",
        "※ 금액 단위: 원. 시가총액은 보통주 + 우선주 합산.",
    ], start=1):
        ws0.cell(row=i, column=1, value=line)
    ws0.column_dimensions["A"].width = 100
    ws0["A1"].font = Font(bold=True, size=13)

    ok = [r for r in results if r.get("status") == "ok"]

    # ── 검증계산 ──
    ws = wb.create_sheet("검증계산")
    cols = ["회사", "종목코드", "기준분기", "연결/별도", "시가총액"]
    col = {}

    def add(name, fill=None):
        cols.append(name)
        col[name] = len(cols)
        return len(cols)

    for k, label in FLOWS:
        add(f"{label} YTD당기", INPUT); add(f"{label} 연간(전년)", INPUT); add(f"{label} YTD전년동기", INPUT)
        add(f"{label} TTM[엑셀]", EXCEL); add(f"{label} TTM[파이썬]"); add(f"{label} 일치")
    for k, label in STOCKS:
        add(f"{label} 당기말", INPUT); add(f"{label} 전년동기말", INPUT)
    for m in ["E/P", "B/P", "ROE", "GP/A", "EV", "EBIT/EV", "CFO/NI"]:
        add(f"{m}[엑셀]", EXCEL); add(f"{m}[파이썬]"); add(f"{m} 일치")
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HDR
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions["A"].width = 18
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"

    L = lambda name: get_column_letter(col[name])  # noqa: E731
    for i, r in enumerate(ok, start=2):
        vi = r["verify_inputs"]
        ws.cell(row=i, column=1, value=r["name"])
        ws.cell(row=i, column=2, value=r["stock_code"])
        ws.cell(row=i, column=3, value=r["period"])
        ws.cell(row=i, column=4, value=r["fs_div"])
        ws.cell(row=i, column=5, value=r["mcap"])
        for k, label in FLOWS:
            a, b, c = (f"{label} YTD당기", f"{label} 연간(전년)", f"{label} YTD전년동기")
            ws.cell(row=i, column=col[a], value=vi[k]["ytd_P"]).fill = INPUT
            ws.cell(row=i, column=col[b], value=vi[k]["annual_prev"]).fill = INPUT
            ws.cell(row=i, column=col[c], value=vi[k]["ytd_P_prevyear"]).fill = INPUT
            ex = f"{label} TTM[엑셀]"
            ws.cell(row=i, column=col[ex],
                    value=f'=IF(OR({L(a)}{i}="",{L(b)}{i}="",{L(c)}{i}=""),"",{L(a)}{i}+{L(b)}{i}-{L(c)}{i})').fill = EXCEL
            ws.cell(row=i, column=col[f"{label} TTM[파이썬]"], value=r["ttm"].get(k))
            ws.cell(row=i, column=col[f"{label} 일치"], value=_match(f"{L(ex)}{i}", f"{L(label + ' TTM[파이썬]')}{i}"))
        for k, label in STOCKS:
            ws.cell(row=i, column=col[f"{label} 당기말"], value=vi[k]["P"]).fill = INPUT
            ws.cell(row=i, column=col[f"{label} 전년동기말"], value=vi[k]["P_prevyear"]).fill = INPUT
        E = lambda n: f"{L(n)}{i}"  # noqa: E731
        mc = f"E{i}"
        f = {
            "E/P": f'=IFERROR({E("지배주주순이익 TTM[엑셀]")}/{mc},"")',
            "B/P": f'=IFERROR(IF({E("지배주주자본 당기말")}<=0,"",{E("지배주주자본 당기말")}/{mc}),"")',
            "ROE": f'=IFERROR(IF(OR({E("지배주주자본 당기말")}<=0,{E("지배주주자본 전년동기말")}<=0),"",'
                   f'{E("지배주주순이익 TTM[엑셀]")}/(({E("지배주주자본 당기말")}+{E("지배주주자본 전년동기말")})/2)),"")',
            "GP/A": f'=IFERROR({E("매출총이익 TTM[엑셀]")}/(({E("자산총계 당기말")}+{E("자산총계 전년동기말")})/2),"")',
            "EV": f'=IFERROR({mc}+{E("이자부부채 당기말")}-{E("현금성자산 당기말")}-{E("단기금융상품 당기말")},"")',
        }
        f["EBIT/EV"] = f'=IFERROR(IF(OR({E("EV[엑셀]")}<=0,{E("EV[엑셀]")}<0.1*{mc}),"",{E("영업이익 TTM[엑셀]")}/{E("EV[엑셀]")}),"")'
        f["CFO/NI"] = f'=IFERROR(IF({E("당기순이익 TTM[엑셀]")}<=0,"",{E("영업현금흐름 TTM[엑셀]")}/{E("당기순이익 TTM[엑셀]")}),"")'
        pyv = {"E/P": r["metrics"]["E/P"], "B/P": r["metrics"]["B/P"], "ROE": r["metrics"]["ROE"], "GP/A": r["metrics"]["GP/A"],
               "EV": r["ev"], "EBIT/EV": r["metrics"]["EBIT/EV"], "CFO/NI": r["metrics"]["CFO/NI"]}
        for m, formula in f.items():
            ws.cell(row=i, column=col[f"{m}[엑셀]"], value=formula).fill = EXCEL
            ws.cell(row=i, column=col[f"{m}[파이썬]"], value=pyv[m])
            ws.cell(row=i, column=col[f"{m} 일치"], value=_match(E(f"{m}[엑셀]"), E(f"{m}[파이썬]")))

    # ── 원천값 확인 ──
    ws2 = wb.create_sheet("원천값 확인")
    hdr = ["회사", "종목코드", "항목", "기간 역할", "금액(원)", "매핑 신뢰도", "원 계정 ID", "원 계정명", "접수번호", "DART 원문", "확인(O/X)", "메모"]
    for c, h in enumerate(hdr, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = HDR
        cell.font = Font(bold=True)
    widths = [16, 9, 14, 14, 18, 12, 40, 26, 16, 14, 9, 30]
    for c, w in enumerate(widths, start=1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = "A2"
    row = 2
    for r in ok:
        rc = r["verify_inputs"]["rcept"]
        for k, label in FLOWS + STOCKS:
            src = (r.get("sources_P") or {}).get(k)
            conf = r["confidence_P"].get(k)
            v = r["verify_inputs"][k]
            entries = ([("당기 YTD", v["ytd_P"], rc["P"]), ("전년 연간", v["annual_prev"], rc["annual_prev"]),
                        ("전년 동기 YTD", v["ytd_P_prevyear"], rc["P_prevyear"])] if "ytd_P" in v
                       else [("당기말", v["P"], rc["P"]), ("전년 동기말", v["P_prevyear"], rc["P_prevyear"])])
            for role, amount, rno in entries:
                ws2.cell(row=row, column=1, value=r["name"])
                ws2.cell(row=row, column=2, value=r["stock_code"])
                ws2.cell(row=row, column=3, value=label)
                ws2.cell(row=row, column=4, value=role)
                ws2.cell(row=row, column=5, value=amount).number_format = "#,##0"
                ws2.cell(row=row, column=6, value=conf if role in ("당기 YTD", "당기말") else "")
                if src and role in ("당기 YTD", "당기말"):
                    ws2.cell(row=row, column=7, value=" + ".join(str(s[0]) for s in src))
                    ws2.cell(row=row, column=8, value=" + ".join(str(s[1]) for s in src))
                ws2.cell(row=row, column=9, value=rno)
                if rno:
                    c = ws2.cell(row=row, column=10, value="열기")
                    c.hyperlink = f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={rno}"
                    c.font = Font(color="1F5FD6", underline="single")
                row += 1

    # ── 커버리지 ──
    ws3 = wb.create_sheet("커버리지")
    ws3.append(["지표", "계산 가능", "평가 대상", "비율", "무효 사유"])
    for c in range(1, 6):
        ws3.cell(row=1, column=c).fill = HDR
        ws3.cell(row=1, column=c).font = Font(bold=True)
    from collections import Counter
    for m in ["E/P", "B/P", "EBIT/EV", "EBITDA/EV", "ROE", "GP/A", "CFO/NI", "이자보상배율", "영업이익증가율", "12-1모멘텀"]:
        vals = [x["metrics"].get(m) for x in ok]
        reasons = Counter(x["metric_reasons"].get(m) or "값 없음" for x in ok if x["metrics"].get(m) is None)
        n = sum(v is not None for v in vals)
        ws3.append([m, n, len(ok), f"{n / max(len(ok), 1):.0%}", ", ".join(f"{k} {v}" for k, v in reasons.items())])
    ws3.append([])
    ws3.append(["평가 제외 회사", "상태", "사유"])
    for x in results:
        if x.get("status") != "ok":
            ws3.append([x.get("name"), x.get("status"), "; ".join(x.get("flags", [])) or x.get("error", "")])
    for c, w in zip("ABCDE", [16, 12, 10, 8, 80]):
        ws3.column_dimensions[c].width = w

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
